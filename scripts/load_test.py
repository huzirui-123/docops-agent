#!/usr/bin/env python3
"""Local load test utility for the FastAPI endpoint.

This script targets a real running service via base_url.
It does not use ASGITransport.
"""

from __future__ import annotations

import argparse
import asyncio
import io
import json
import math
import time
from collections import Counter
from dataclasses import dataclass
from typing import Any

import httpx
from docx import Document


@dataclass(frozen=True)
class RequestResult:
    status_code: int
    latency_ms: int
    request_id: str | None


def _build_docx_bytes(skill: str) -> bytes:
    placeholders = {
        "meeting_notice": "【MEETING_TITLE】",
        "training_notice": "【TRAINING_TITLE】",
        "inspection_record": "【INSPECTION_SUBJECT】",
    }
    placeholder = placeholders.get(skill)
    if placeholder is None:
        raise ValueError(f"Unsupported skill: {skill}")

    document = Document()
    document.add_paragraph(placeholder)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_task_bytes(skill: str) -> bytes:
    payloads: dict[str, dict[str, Any]] = {
        "meeting_notice": {
            "task_type": "meeting_notice",
            "payload": {"meeting_title": "LoadTest Meeting"},
        },
        "training_notice": {
            "task_type": "training_notice",
            "payload": {"training_title": "LoadTest Training"},
        },
        "inspection_record": {
            "task_type": "inspection_record",
            "payload": {"inspection_subject": "LoadTest Site"},
        },
    }
    task = payloads.get(skill)
    if task is None:
        raise ValueError(f"Unsupported skill: {skill}")
    return json.dumps(task, ensure_ascii=False).encode("utf-8")


def _percentile(values: list[int], p: float) -> int:
    if not values:
        return 0
    ordered = sorted(values)
    index = min(len(ordered) - 1, max(0, math.ceil((p / 100) * len(ordered)) - 1))
    return ordered[index]


async def _run_single(
    client: httpx.AsyncClient,
    *,
    base_url: str,
    skill: str,
    template_bytes: bytes,
    task_bytes: bytes,
) -> RequestResult:
    started = time.perf_counter()
    request_id: str | None = None
    status_code = 599

    try:
        response = await client.post(
            f"{base_url.rstrip('/')}/v1/run",
            files={
                "template": (
                    "template.docx",
                    template_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ),
                "task": ("task.json", task_bytes, "application/json"),
            },
            data={"skill": skill},
        )
        status_code = response.status_code
        request_id = response.headers.get("X-Docops-Request-Id")
        if status_code == 408 and request_id is None:
            try:
                payload = response.json()
                detail = payload.get("detail", {}) if isinstance(payload, dict) else {}
                request_id = detail.get("request_id")
            except Exception:  # noqa: BLE001
                request_id = None
    except Exception:  # noqa: BLE001
        status_code = 599

    return RequestResult(
        status_code=status_code,
        latency_ms=int((time.perf_counter() - started) * 1000),
        request_id=request_id,
    )


async def _run_load_test(
    *,
    base_url: str,
    skill: str,
    concurrency: int,
    requests: int,
    timeout: float,
) -> dict[str, Any]:
    semaphore = asyncio.Semaphore(concurrency)
    template_bytes = _build_docx_bytes(skill)
    task_bytes = _build_task_bytes(skill)

    async with httpx.AsyncClient(timeout=timeout) as client:
        async def wrapped() -> RequestResult:
            async with semaphore:
                return await _run_single(
                    client,
                    base_url=base_url,
                    skill=skill,
                    template_bytes=template_bytes,
                    task_bytes=task_bytes,
                )

        tasks = [asyncio.create_task(wrapped()) for _ in range(requests)]
        results = await asyncio.gather(*tasks)

    status_counts = Counter(str(item.status_code) for item in results)
    latencies = [item.latency_ms for item in results]
    timeout_request_ids = [item.request_id for item in results if item.status_code == 408]

    return {
        "base_url": base_url,
        "skill": skill,
        "requests": requests,
        "concurrency": concurrency,
        "status_counts": dict(sorted(status_counts.items())),
        "latency_ms": {
            "p50": _percentile(latencies, 50),
            "p95": _percentile(latencies, 95),
            "avg": int(sum(latencies) / len(latencies)) if latencies else 0,
        },
        "timeout_request_ids": [rid for rid in timeout_request_ids if rid],
        "note": "Run against a real server process (uvicorn/gunicorn).",
    }


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Local load test for docops /v1/run")
    parser.add_argument("--base-url", default="http://127.0.0.1:8000")
    parser.add_argument("--skill", default="meeting_notice")
    parser.add_argument("--concurrency", type=int, default=8)
    parser.add_argument("--requests", type=int, default=20)
    parser.add_argument("--timeout", type=float, default=30.0)
    return parser.parse_args()


def main() -> None:
    args = _parse_args()
    summary = asyncio.run(
        _run_load_test(
            base_url=args.base_url,
            skill=args.skill,
            concurrency=args.concurrency,
            requests=args.requests,
            timeout=args.timeout,
        )
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
