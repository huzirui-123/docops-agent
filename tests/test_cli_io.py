from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from apps.cli.io import build_output_paths, write_render_output_atomic
from core.format.models import FormatReport
from core.render.models import MissingFieldsReport, RenderOutput, ReplaceReport, ReplaceSummary
from core.templates.models import ParseResult


def _build_output() -> RenderOutput:
    document = Document()
    document.add_paragraph("hello")

    return RenderOutput(
        document=document,
        parse_result=ParseResult(),
        template_fields=[],
        replace_report=ReplaceReport(
            entries=[],
            summary=ReplaceSummary(
                total_placeholders=0,
                replaced_count=0,
                missing_count=0,
                had_unsupported=False,
                unsupported_count=0,
                unsupported_mode="error",
            ),
            touched_runs=[],
        ),
        missing_fields=MissingFieldsReport(missing_required=[], missing_optional=[]),
        format_report=FormatReport(passed=True, error_count=0, fixed_count=0, issues=[]),
    )


def test_write_render_output_atomic_cleans_docx_tmp_on_success(tmp_path: Path) -> None:
    output = _build_output()
    paths = build_output_paths(tmp_path)

    write_render_output_atomic(paths, output)

    assert paths.docx.exists()
    assert list(tmp_path.glob("out.docx.*.tmp")) == []


def test_write_render_output_atomic_cleans_docx_tmp_on_failure(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    output = _build_output()
    paths = build_output_paths(tmp_path)

    def broken_save(_: str) -> None:
        raise RuntimeError("save failed")

    monkeypatch.setattr(output.document, "save", broken_save)

    with pytest.raises(RuntimeError, match="save failed"):
        write_render_output_atomic(paths, output)

    assert not paths.docx.exists()
    assert list(tmp_path.glob("out.docx.*.tmp")) == []

