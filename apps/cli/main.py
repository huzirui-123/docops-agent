"""Typer CLI entrypoint for docops-agent."""

from __future__ import annotations

import json
from collections import Counter
from pathlib import Path
from typing import Annotated, Any, Literal, cast

import typer
from docx import Document

from apps.cli.format_human import render_format_summary
from apps.cli.io import (
    OutputPaths,
    build_debug_output_path,
    build_output_paths,
    existing_output_files,
    write_debug_dump_atomic,
    write_fallback_json_atomic,
    write_render_output_atomic,
    write_suggested_policy_atomic,
)
from core.format.models import FormatIssue
from core.format.policy_loader import load_policy
from core.format.suggested_policy import build_suggested_policy
from core.orchestrator.pipeline import run_task
from core.render.debug_dump import DebugReport, collect_suspicious_runs
from core.render.models import RenderOutput
from core.skills.base import Skill
from core.skills.meeting_notice import MeetingNoticeSkill
from core.skills.models import TaskSpec
from core.utils.errors import MissingRequiredFieldsError, TemplateError

app = typer.Typer(help="Document Ops Agent CLI", rich_markup_mode=None)
UnsupportedMode = Literal["error", "warn"]
FormatMode = Literal["report", "strict", "off"]
FormatBaseline = Literal["template", "policy"]
FormatFixMode = Literal["none", "safe"]
FormatReportMode = Literal["human", "json", "both"]


@app.callback()
def cli_callback() -> None:
    """CLI root callback to keep `docops run` as explicit command form."""


@app.command("run")
def run_command(
    template: Annotated[Path, typer.Option(..., exists=True, dir_okay=False, file_okay=True)],
    task: Annotated[Path, typer.Option(..., exists=True, dir_okay=False, file_okay=True)],
    skill: Annotated[str, typer.Option(...)],
    out_dir: Annotated[Path, typer.Option()] = Path("."),
    policy: Annotated[Path | None, typer.Option()] = None,
    unsupported_mode: Annotated[str, typer.Option()] = "error",
    format_mode: Annotated[str, typer.Option()] = "report",
    format_baseline: Annotated[str, typer.Option()] = "template",
    format_fix_mode: Annotated[str, typer.Option()] = "safe",
    format_report: Annotated[str, typer.Option()] = "human",
    export_suggested_policy: Annotated[
        Path | None,
        typer.Option(
            "--export-suggested-policy",
            help="Write optional suggested policy YAML for template baseline tuning.",
        ),
    ] = None,
    force: Annotated[
        bool, typer.Option("--force", help="Overwrite outputs when they already exist.")
    ] = False,
    no_overwrite: Annotated[
        bool,
        typer.Option(
            "--no-overwrite",
            help="Fail when outputs already exist.",
        ),
    ] = False,
    debug_dump: Annotated[
        bool,
        typer.Option(
            "--debug-dump",
            help="Write out.debug.json with suspicious symbol/font diagnostics.",
        ),
    ] = False,
) -> None:
    """Execute one document generation run and write fixed output artifacts."""

    paths = build_output_paths(out_dir)
    explicit_overwrite = True
    normalized_mode = unsupported_mode.lower().strip()
    if normalized_mode not in {"error", "warn"}:
        typer.echo("ERROR: --unsupported-mode must be one of: error, warn.")
        _safe_write_exit1_fallback(
            paths,
            "ArgumentValidationError",
            "invalid unsupported_mode",
            "args",
        )
        raise typer.Exit(code=1)
    unsupported_mode_typed = cast(UnsupportedMode, normalized_mode)

    normalized_format_mode = format_mode.lower().strip()
    if normalized_format_mode not in {"report", "strict", "off"}:
        typer.echo("ERROR: --format-mode must be one of: report, strict, off.")
        _safe_write_exit1_fallback(
            paths,
            "ArgumentValidationError",
            "invalid format_mode",
            "args",
        )
        raise typer.Exit(code=1)
    format_mode_typed = cast(FormatMode, normalized_format_mode)

    normalized_format_baseline = format_baseline.lower().strip()
    if normalized_format_baseline not in {"template", "policy"}:
        typer.echo("ERROR: --format-baseline must be one of: template, policy.")
        _safe_write_exit1_fallback(
            paths,
            "ArgumentValidationError",
            "invalid format_baseline",
            "args",
        )
        raise typer.Exit(code=1)
    format_baseline_typed = cast(FormatBaseline, normalized_format_baseline)

    normalized_format_fix_mode = format_fix_mode.lower().strip()
    if normalized_format_fix_mode not in {"none", "safe"}:
        typer.echo("ERROR: --format-fix-mode must be one of: none, safe.")
        _safe_write_exit1_fallback(
            paths,
            "ArgumentValidationError",
            "invalid format_fix_mode",
            "args",
        )
        raise typer.Exit(code=1)
    format_fix_mode_typed = cast(FormatFixMode, normalized_format_fix_mode)

    normalized_format_report = format_report.lower().strip()
    if normalized_format_report not in {"human", "json", "both"}:
        typer.echo("ERROR: --format-report must be one of: human, json, both.")
        _safe_write_exit1_fallback(
            paths,
            "ArgumentValidationError",
            "invalid format_report",
            "args",
        )
        raise typer.Exit(code=1)
    format_report_typed = cast(FormatReportMode, normalized_format_report)

    if force and no_overwrite:
        typer.echo("ERROR: --force and --no-overwrite cannot be used together.")
        _safe_write_exit1_fallback(paths, "ArgumentConflict", "conflicting overwrite flags", "args")
        raise typer.Exit(code=1)

    if no_overwrite:
        explicit_overwrite = False
    elif force:
        explicit_overwrite = True

    extra_outputs: list[Path] = []
    if export_suggested_policy is not None:
        extra_outputs.append(export_suggested_policy)

    existing = existing_output_files(paths, extra_paths=extra_outputs)
    if existing and explicit_overwrite:
        names = ", ".join(path.name for path in existing)
        typer.echo(f"INFO: overwriting existing outputs: {names}")

    if existing and not explicit_overwrite:
        typer.echo("ERROR: outputs already exist and --no-overwrite is enabled.")
        raise typer.Exit(code=1)

    typer.echo(f"INFO(format): baseline={format_baseline_typed}")

    output: RenderOutput | None = None
    template_document = None
    template_for_suggested_policy = None
    policy_model = None
    debug_pre_report: DebugReport | None = None
    exit_code = 1
    reason = "unexpected error"
    generic_error: Exception | None = None
    failure_stage = "unknown"

    try:
        failure_stage = "load_task"
        task_spec = _load_task_spec(task)
        failure_stage = "resolve_skill"
        selected_skill = _resolve_skill(skill)
        failure_stage = "load_policy"
        policy_model = load_policy(policy)
        failure_stage = "load_template"
        document = Document(str(template))
        template_document = document
        if export_suggested_policy is not None:
            template_for_suggested_policy = Document(str(template))
        if debug_dump:
            debug_pre_report = collect_suspicious_runs(
                document=document,
                touched_runs=set(),
                stage="pre_pipeline",
            )
        failure_stage = "pipeline"

        output = run_task(
            task_spec=task_spec,
            template_document=document,
            skill=selected_skill,
            policy=policy_model,
            unsupported_mode=unsupported_mode_typed,
            format_mode=format_mode_typed,
            format_baseline=format_baseline_typed,
            format_fix_mode=format_fix_mode_typed,
        )

        if output.replace_report.summary.had_unsupported and unsupported_mode_typed == "warn":
            typer.echo(
                "WARNING(unsupported): unsupported placeholders detected "
                f"(count={output.replace_report.summary.unsupported_count}, mode=warn)."
            )

        if format_mode_typed == "off":
            typer.echo("INFO(format): fix skipped (off mode)")
            typer.echo("INFO: format validation skipped")
            exit_code = 0
            reason = "success"
        elif format_mode_typed == "strict" and not output.format_report.passed:
            exit_code = 4
            reason = "format validation failed"
        else:
            if output.format_report.issues:
                typer.echo(
                    "WARNING(format): format issues detected "
                    f"({_format_issue_summary(output.format_report.issues)})."
                )
            exit_code = 0
            reason = "success"

    except TemplateError as exc:
        output = exc.render_output
        exit_code = 3
        reason = "template unsupported placeholders"
        typer.echo(f"ERROR: {reason}")
    except MissingRequiredFieldsError as exc:
        output = exc.render_output
        exit_code = 2
        reason = "missing required fields"
        typer.echo(f"ERROR: {reason}")
        if (
            output is not None
            and output.replace_report.summary.had_unsupported
            and unsupported_mode_typed == "warn"
        ):
            typer.echo(
                "WARNING(unsupported): unsupported placeholders detected "
                f"(count={output.replace_report.summary.unsupported_count}, mode=warn)."
            )
    except Exception as exc:  # noqa: BLE001
        generic_error = exc
        exit_code = 1
        reason = "internal error"
        typer.echo(f"ERROR: {type(exc).__name__}: {exc}")

    docx_write_error: str | None = None

    if (
        output is not None
        and output.format_report.summary is not None
        and format_report_typed in {"human", "both"}
    ):
        typer.echo(render_format_summary(output, format_fix_mode_typed))

    if output is not None:
        try:
            write_render_output_atomic(paths, output)
        except Exception as write_exc:  # noqa: BLE001
            docx_write_error = str(write_exc)
            exit_code = 1
            reason = "write output failed"
            typer.echo(f"ERROR: {reason}: {docx_write_error}")
            _safe_write_exit1_fallback(
                paths,
                error_type=type(write_exc).__name__,
                error_message=str(write_exc),
                stage="write_docx",
                docx_write_error=docx_write_error,
                base_output=output,
            )

        if (
            export_suggested_policy is not None
            and policy_model is not None
            and docx_write_error is None
        ):
            try:
                source_document = (
                    template_for_suggested_policy
                    if template_for_suggested_policy is not None
                    else output.document
                )
                suggested_policy = build_suggested_policy(source_document, policy_model)
                write_suggested_policy_atomic(export_suggested_policy, suggested_policy)
                typer.echo(f"INFO: wrote suggested policy to {export_suggested_policy}")
            except Exception as exc:  # noqa: BLE001
                exit_code = 1
                reason = "write suggested policy failed"
                typer.echo(f"ERROR: {reason}: {exc}")
    else:
        error_type = type(generic_error).__name__ if generic_error is not None else "UnknownError"
        error_message = str(generic_error) if generic_error is not None else reason
        _safe_write_exit1_fallback(
            paths,
            error_type=error_type,
            error_message=error_message,
            stage=failure_stage,
            docx_write_error=docx_write_error,
        )

    if debug_dump:
        debug_payload = _build_debug_payload(
            pre_report=debug_pre_report,
            output=output,
            template_document=template_document,
        )
        if debug_payload is not None:
            debug_path = build_debug_output_path(out_dir)
            try:
                write_debug_dump_atomic(debug_path, debug_payload)
            except Exception as exc:  # noqa: BLE001
                typer.echo(f"ERROR: debug dump write failed: {exc}")
                exit_code = 1

    if exit_code == 0:
        typer.echo("INFO: success")
    elif exit_code == 4:
        typer.echo("ERROR: format validation failed")

    raise typer.Exit(code=exit_code)


def _load_task_spec(path: Path) -> TaskSpec:
    raw = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(raw, dict):
        raise ValueError("Task JSON must be an object")
    return TaskSpec.model_validate(raw)


def _resolve_skill(skill_name: str) -> Skill:
    if skill_name == "meeting_notice":
        return MeetingNoticeSkill()
    raise ValueError(f"Unsupported skill: {skill_name}")


def _safe_write_exit1_fallback(
    paths: OutputPaths,
    error_type: str,
    error_message: str,
    stage: str,
    docx_write_error: str | None = None,
    base_output: RenderOutput | None = None,
) -> None:
    try:
        write_fallback_json_atomic(
            paths,
            error_type=error_type,
            error_message=error_message,
            stage=stage,
            docx_write_error=docx_write_error,
            base_output=base_output,
        )
    except Exception:  # noqa: BLE001
        pass


def _build_debug_payload(
    *,
    pre_report: DebugReport | None,
    output: RenderOutput | None,
    template_document: Any | None,
) -> dict[str, Any] | None:
    if output is not None:
        touched_runs = set(output.replace_report.touched_runs)
        post_report = collect_suspicious_runs(
            document=output.document,
            touched_runs=touched_runs,
            stage="post_pipeline",
        )
    elif template_document is not None:
        post_report = collect_suspicious_runs(
            document=template_document,
            touched_runs=set(),
            stage="post_pipeline",
        )
    else:
        return None

    payload: dict[str, Any] = {
        "post_pipeline": post_report.model_dump(mode="json"),
    }
    if pre_report is not None:
        payload["pre_pipeline"] = pre_report.model_dump(mode="json")
    return payload


def _format_issue_summary(issues: list[FormatIssue]) -> str:
    counter: Counter[str] = Counter(issue.code for issue in issues)
    return ", ".join(f"{code}={counter[code]}" for code in sorted(counter))


def main() -> None:
    """Poetry script entrypoint."""

    app()


if __name__ == "__main__":
    main()
