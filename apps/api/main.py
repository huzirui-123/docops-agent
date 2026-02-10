"""FastAPI wrapper for docops generation pipeline (M7-1)."""

from __future__ import annotations

import importlib.metadata
import json
import os
import queue
import shutil
import tempfile
import threading
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Annotated, Any, Literal, cast

from docx import Document
from fastapi import BackgroundTasks, FastAPI, File, Form, UploadFile
from fastapi.responses import JSONResponse, Response
from pydantic import ValidationError

from apps.cli.io import (
    build_output_paths,
    write_render_output_atomic,
    write_suggested_policy_atomic,
)
from core.format.policy_loader import load_policy
from core.format.suggested_policy import build_suggested_policy
from core.orchestrator.pipeline import run_task
from core.skills.base import Skill
from core.skills.meeting_notice import MeetingNoticeSkill
from core.skills.models import TaskSpec
from core.utils.errors import MissingRequiredFieldsError, TemplateError

app = FastAPI(title="docops-agent API", version="0.1.0")

FormatMode = Literal["report", "strict", "off"]
FormatBaseline = Literal["template", "policy"]
FormatFixMode = Literal["none", "safe"]
FormatReportMode = Literal["human", "json", "both"]
PresetMode = Literal["quick", "template", "strict"]

_DEFAULT_MAX_UPLOAD_BYTES = 25 * 1024 * 1024
_DEFAULT_TIMEOUT_SECONDS = 60.0

_PRESET_TO_FORMAT: dict[PresetMode, tuple[FormatMode, FormatBaseline, FormatFixMode]] = {
    "quick": ("report", "template", "safe"),
    "template": ("report", "template", "none"),
    "strict": ("strict", "policy", "safe"),
}


@dataclass(frozen=True)
class EffectiveConfig:
    preset: PresetMode
    format_mode: FormatMode
    format_baseline: FormatBaseline
    format_fix_mode: FormatFixMode
    format_report: FormatReportMode


@dataclass(frozen=True)
class RunExecutionResult:
    exit_code: int
    message: str


class ApiRequestError(Exception):
    def __init__(
        self,
        *,
        status_code: int,
        error_code: str,
        message: str,
        detail: dict[str, Any] | None = None,
    ) -> None:
        super().__init__(message)
        self.status_code = status_code
        self.error_code = error_code
        self.message = message
        self.detail = detail or {}


@app.get("/healthz")
async def healthz() -> dict[str, str]:
    """Liveness endpoint."""

    return {"status": "ok"}


@app.post("/v1/run", response_model=None)
async def run_v1(
    template: Annotated[UploadFile, File(...)],
    task: Annotated[UploadFile, File(...)],
    skill: Annotated[str, Form()] = "meeting_notice",
    preset: Annotated[str | None, Form()] = None,
    format_mode: Annotated[str | None, Form()] = None,
    format_baseline: Annotated[str | None, Form()] = None,
    format_fix_mode: Annotated[str | None, Form()] = None,
    format_report: Annotated[str | None, Form()] = None,
    policy_yaml: Annotated[str | None, Form()] = None,
    export_suggested_policy: Annotated[bool, Form()] = False,
) -> Response:
    """Run one generation job and return a zip with output artifacts."""

    tmp_dir = Path(tempfile.mkdtemp(prefix="docops-api-run-"))
    zip_path: Path | None = None

    try:
        max_upload_bytes = _max_upload_bytes()
        timeout_seconds = _timeout_seconds()

        _validate_upload_name(template.filename, expected_suffix=".docx", field_name="template")
        _validate_upload_name(task.filename, expected_suffix=".json", field_name="task")

        template_path = tmp_dir / "template.docx"
        task_path = tmp_dir / "task.json"

        template_magic = _save_upload_with_limit(
            upload=template,
            destination=template_path,
            max_bytes=max_upload_bytes,
            field_name="template",
        )
        if template_magic != b"PK\x03\x04":
            raise ApiRequestError(
                status_code=415,
                error_code="INVALID_MEDIA_TYPE",
                message="template must be a valid .docx file",
                detail={"field": "template"},
            )

        _save_upload_with_limit(
            upload=task,
            destination=task_path,
            max_bytes=max_upload_bytes,
            field_name="task",
        )

        task_spec = _load_task_spec(task_path)
        selected_skill = _resolve_skill(skill)

        effective = _resolve_effective_config(
            preset_input=preset,
            format_mode_input=format_mode,
            format_baseline_input=format_baseline,
            format_fix_mode_input=format_fix_mode,
            format_report_input=format_report,
            policy_yaml_input=policy_yaml,
        )

        policy_path: Path | None = None
        if policy_yaml is not None:
            policy_path = tmp_dir / "policy.yaml"
            policy_path.write_text(policy_yaml, encoding="utf-8")

        policy_model = _load_policy_with_api_error(policy_path)

        run_result = _run_pipeline_with_timeout(
            timeout_seconds=timeout_seconds,
            task_spec=task_spec,
            template_path=template_path,
            selected_skill=selected_skill,
            policy_model=policy_model,
            effective=effective,
        )

        if run_result.output is None:
            raise RuntimeError("Pipeline did not produce render output")

        paths = build_output_paths(tmp_dir)
        write_render_output_atomic(paths, run_result.output)

        suggested_policy_path: Path | None = None
        if export_suggested_policy:
            suggested_policy_path = tmp_dir / "out.suggested_policy.yaml"
            template_for_suggested = Document(str(template_path))
            source_doc = (
                template_for_suggested
                if template_for_suggested is not None
                else run_result.output.document
            )
            suggested_payload = build_suggested_policy(source_doc, policy_model)
            write_suggested_policy_atomic(suggested_policy_path, suggested_payload)

        api_result_path = tmp_dir / "api_result.json"
        _write_json_atomic(
            api_result_path,
            _build_api_result(
                exit_code=run_result.exit_code,
                message=run_result.message,
                input_payload={
                    "skill": skill,
                    "preset": preset,
                    "format_mode": format_mode,
                    "format_baseline": format_baseline,
                    "format_fix_mode": format_fix_mode,
                    "format_report": format_report,
                    "policy_yaml_provided": policy_yaml is not None,
                    "export_suggested_policy": export_suggested_policy,
                },
                effective=effective,
            ),
        )

        required_paths = [
            paths.docx,
            paths.replace_log,
            paths.missing_fields,
            paths.format_report,
            api_result_path,
        ]
        for required in required_paths:
            if not required.exists():
                raise RuntimeError(f"Missing output artifact: {required.name}")

        optional_paths: list[Path] = []
        if suggested_policy_path is not None:
            optional_paths.append(suggested_policy_path)

        zip_path = _create_zip(required_paths, optional_paths)

        headers = {"X-Docops-Exit-Code": str(run_result.exit_code)}
        cleanup_tasks = BackgroundTasks()
        cleanup_tasks.add_task(_cleanup_after_response, zip_path, tmp_dir)

        zip_bytes = zip_path.read_bytes()
        headers["Content-Disposition"] = 'attachment; filename=\"docops_outputs.zip\"'
        return Response(
            content=zip_bytes,
            media_type="application/zip",
            headers=headers,
            background=cleanup_tasks,
        )
    except ApiRequestError as exc:
        _cleanup_now(zip_path, tmp_dir)
        return _error_response(
            status_code=exc.status_code,
            error_code=exc.error_code,
            message=exc.message,
            detail=exc.detail,
        )
    except TimeoutError:
        _cleanup_now(zip_path, tmp_dir)
        return _error_response(
            status_code=408,
            error_code="REQUEST_TIMEOUT",
            message="request timed out",
            detail={"timeout_seconds": _timeout_seconds()},
        )
    except Exception as exc:  # noqa: BLE001
        _cleanup_now(zip_path, tmp_dir)
        return _error_response(
            status_code=500,
            error_code="INTERNAL_ERROR",
            message="internal server error",
            detail={"error": str(exc)},
        )


@dataclass
class _PipelineResult:
    output: Any | None
    exit_code: int
    message: str


def _run_pipeline_with_timeout(
    *,
    timeout_seconds: float,
    task_spec: TaskSpec,
    template_path: Path,
    selected_skill: Skill,
    policy_model,
    effective: EffectiveConfig,
) -> _PipelineResult:
    """Execute pipeline with timeout using an isolated thread.

    Keep timeout behavior explicit without depending on asyncio thread helpers.
    This helper keeps timeout behavior explicit and deterministic.
    """

    outcome: queue.Queue[tuple[str, object]] = queue.Queue(maxsize=1)

    def _runner() -> None:
        try:
            result = _run_pipeline_sync(
                task_spec,
                template_path,
                selected_skill,
                policy_model,
                effective,
            )
        except Exception as exc:  # noqa: BLE001
            outcome.put(("error", exc))
            return
        outcome.put(("ok", result))

    worker = threading.Thread(target=_runner, daemon=True, name="docops-api-run")
    worker.start()
    worker.join(timeout_seconds)

    if worker.is_alive():
        raise TimeoutError

    try:
        status, payload = outcome.get_nowait()
    except queue.Empty as exc:
        raise RuntimeError("Pipeline worker ended without producing a result") from exc

    if status == "error":
        error = cast(Exception, payload)
        raise error

    return cast(_PipelineResult, payload)


def _run_pipeline_sync(
    task_spec: TaskSpec,
    template_path: Path,
    selected_skill: Skill,
    policy_model,
    effective: EffectiveConfig,
) -> _PipelineResult:
    template_document = Document(str(template_path))

    try:
        output = run_task(
            task_spec=task_spec,
            template_document=template_document,
            skill=selected_skill,
            policy=policy_model,
            unsupported_mode="error",
            format_mode=effective.format_mode,
            format_baseline=effective.format_baseline,
            format_fix_mode=effective.format_fix_mode,
        )
    except TemplateError as exc:
        return _PipelineResult(
            output=exc.render_output,
            exit_code=3,
            message="template unsupported placeholders",
        )
    except MissingRequiredFieldsError as exc:
        return _PipelineResult(
            output=exc.render_output,
            exit_code=2,
            message="missing required fields",
        )

    if effective.format_mode == "strict" and not output.format_report.passed:
        return _PipelineResult(output=output, exit_code=4, message="format validation failed")

    return _PipelineResult(output=output, exit_code=0, message="success")


def _resolve_effective_config(
    *,
    preset_input: str | None,
    format_mode_input: str | None,
    format_baseline_input: str | None,
    format_fix_mode_input: str | None,
    format_report_input: str | None,
    policy_yaml_input: str | None,
) -> EffectiveConfig:
    advanced_explicit = any(
        value is not None
        for value in (
            format_mode_input,
            format_baseline_input,
            format_fix_mode_input,
            format_report_input,
            policy_yaml_input,
        )
    )

    if preset_input is not None and advanced_explicit:
        raise ApiRequestError(
            status_code=400,
            error_code="INVALID_ARGUMENT_CONFLICT",
            message="preset cannot be combined with advanced format arguments",
            detail={
                "conflicting_fields": [
                    name
                    for name, value in (
                        ("format_mode", format_mode_input),
                        ("format_baseline", format_baseline_input),
                        ("format_fix_mode", format_fix_mode_input),
                        ("format_report", format_report_input),
                        ("policy_yaml", policy_yaml_input),
                    )
                    if value is not None
                ]
            },
        )

    preset_value = _normalize_or_default(
        value=preset_input,
        default="quick",
        field_name="preset",
        allowed={"quick", "template", "strict"},
    )
    preset = cast(PresetMode, preset_value)
    base_mode, base_baseline, base_fix_mode = _PRESET_TO_FORMAT[preset]

    mode = cast(
        FormatMode,
        _normalize_or_default(
            value=format_mode_input,
            default=base_mode,
            field_name="format_mode",
            allowed={"report", "strict", "off"},
        ),
    )
    baseline = cast(
        FormatBaseline,
        _normalize_or_default(
            value=format_baseline_input,
            default=base_baseline,
            field_name="format_baseline",
            allowed={"template", "policy"},
        ),
    )
    fix_mode = cast(
        FormatFixMode,
        _normalize_or_default(
            value=format_fix_mode_input,
            default=base_fix_mode,
            field_name="format_fix_mode",
            allowed={"none", "safe"},
        ),
    )

    # API defaults to quiet mode unless explicitly provided.
    report_mode = cast(
        FormatReportMode,
        _normalize_or_default(
            value=format_report_input,
            default="json",
            field_name="format_report",
            allowed={"human", "json", "both"},
        ),
    )

    return EffectiveConfig(
        preset=preset,
        format_mode=mode,
        format_baseline=baseline,
        format_fix_mode=fix_mode,
        format_report=report_mode,
    )


def _normalize_or_default(
    *,
    value: str | None,
    default: str,
    field_name: str,
    allowed: set[str],
) -> str:
    if value is None:
        return default

    normalized = value.strip().lower()
    if normalized not in allowed:
        raise ApiRequestError(
            status_code=400,
            error_code="INVALID_ARGUMENT",
            message=f"invalid {field_name}",
            detail={"field": field_name, "value": value, "allowed": sorted(allowed)},
        )
    return normalized


def _load_task_spec(task_path: Path) -> TaskSpec:
    try:
        task_raw = json.loads(task_path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ApiRequestError(
            status_code=400,
            error_code="INVALID_JSON",
            message="task file must be valid JSON",
            detail={"field": "task", "error": str(exc)},
        ) from exc
    except UnicodeDecodeError as exc:
        raise ApiRequestError(
            status_code=400,
            error_code="INVALID_JSON",
            message="task file must be UTF-8 JSON",
            detail={"field": "task"},
        ) from exc

    if not isinstance(task_raw, dict):
        raise ApiRequestError(
            status_code=400,
            error_code="INVALID_JSON",
            message="task JSON must be an object",
            detail={"field": "task"},
        )

    try:
        return TaskSpec.model_validate(task_raw)
    except ValidationError as exc:
        raise ApiRequestError(
            status_code=400,
            error_code="INVALID_ARGUMENT",
            message="task JSON schema validation failed",
            detail={"field": "task", "error": str(exc)},
        ) from exc


def _resolve_skill(skill_name: str) -> Skill:
    if skill_name == "meeting_notice":
        return MeetingNoticeSkill()
    raise ApiRequestError(
        status_code=400,
        error_code="INVALID_ARGUMENT",
        message="unsupported skill",
        detail={"field": "skill", "value": skill_name},
    )


def _load_policy_with_api_error(policy_path: Path | None):
    try:
        return load_policy(policy_path)
    except ValueError as exc:
        raise ApiRequestError(
            status_code=400,
            error_code="INVALID_ARGUMENT",
            message="invalid policy_yaml",
            detail={"field": "policy_yaml", "error": str(exc)},
        ) from exc


def _validate_upload_name(filename: str | None, *, expected_suffix: str, field_name: str) -> None:
    if filename is None or not filename.lower().endswith(expected_suffix):
        raise ApiRequestError(
            status_code=415,
            error_code="INVALID_MEDIA_TYPE",
            message=f"{field_name} must be a {expected_suffix} file",
            detail={"field": field_name, "filename": filename},
        )


def _save_upload_with_limit(
    *,
    upload: UploadFile,
    destination: Path,
    max_bytes: int,
    field_name: str,
) -> bytes:
    destination.parent.mkdir(parents=True, exist_ok=True)

    total_size = 0
    magic = b""

    source = upload.file
    source.seek(0)

    with destination.open("wb") as handle:
        while True:
            chunk = source.read(1024 * 1024)
            if not chunk:
                break
            if len(magic) < 4:
                magic = (magic + chunk)[:4]

            total_size += len(chunk)
            if total_size > max_bytes:
                raise ApiRequestError(
                    status_code=413,
                    error_code="UPLOAD_TOO_LARGE",
                    message=f"{field_name} exceeds upload size limit",
                    detail={
                        "field": field_name,
                        "max_bytes": max_bytes,
                        "received_bytes": total_size,
                    },
                )
            handle.write(chunk)

    source.close()
    return magic


def _max_upload_bytes() -> int:
    raw = os.getenv("DOCOPS_MAX_UPLOAD_BYTES")
    if raw is None:
        return _DEFAULT_MAX_UPLOAD_BYTES
    try:
        parsed = int(raw)
    except ValueError:
        return _DEFAULT_MAX_UPLOAD_BYTES
    return parsed if parsed > 0 else _DEFAULT_MAX_UPLOAD_BYTES


def _timeout_seconds() -> float:
    raw = os.getenv("DOCOPS_REQUEST_TIMEOUT_SECONDS")
    if raw is None:
        return _DEFAULT_TIMEOUT_SECONDS
    try:
        parsed = float(raw)
    except ValueError:
        return _DEFAULT_TIMEOUT_SECONDS
    return parsed if parsed > 0 else _DEFAULT_TIMEOUT_SECONDS


def _write_json_atomic(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        mode="w",
        encoding="utf-8",
        dir=path.parent,
        delete=False,
        prefix=f"{path.name}.",
        suffix=".tmp",
    ) as tmp:
        tmp_path = Path(tmp.name)
        json.dump(payload, tmp, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    tmp_path.replace(path)


def _create_zip(required_paths: list[Path], optional_paths: list[Path]) -> Path:
    fd, raw_zip = tempfile.mkstemp(prefix="docops-api-zip-", suffix=".zip")
    os.close(fd)
    zip_path = Path(raw_zip)

    with zipfile.ZipFile(zip_path, mode="w", compression=zipfile.ZIP_DEFLATED) as archive:
        for artifact in required_paths:
            archive.write(artifact, arcname=artifact.name)
        for artifact in optional_paths:
            if artifact.exists():
                archive.write(artifact, arcname=artifact.name)

    return zip_path


def _build_api_result(
    *,
    exit_code: int,
    message: str,
    input_payload: dict[str, Any],
    effective: EffectiveConfig,
) -> dict[str, Any]:
    return {
        "exit_code": exit_code,
        "message": message,
        "input": input_payload,
        "effective": {
            "preset": effective.preset,
            "format_mode": effective.format_mode,
            "format_baseline": effective.format_baseline,
            "format_fix_mode": effective.format_fix_mode,
            "format_report": effective.format_report,
        },
        "build": {
            "version": _package_version(),
            "commit": os.getenv("DOCOPS_COMMIT_SHA", "unknown"),
        },
    }


def _package_version() -> str:
    try:
        return importlib.metadata.version("docops-agent")
    except importlib.metadata.PackageNotFoundError:
        return "unknown"


def _error_response(
    *,
    status_code: int,
    error_code: str,
    message: str,
    detail: dict[str, Any] | None = None,
) -> JSONResponse:
    return JSONResponse(
        status_code=status_code,
        content={
            "error_code": error_code,
            "message": message,
            "detail": detail or {},
        },
    )


def _cleanup_now(zip_path: Path | None, tmp_dir: Path) -> None:
    if zip_path is not None:
        _safe_remove_file(zip_path)
    _safe_remove_dir(tmp_dir)


async def _cleanup_after_response(zip_path: Path, tmp_dir: Path) -> None:
    """Cleanup generated temp artifacts after response is sent."""

    _safe_remove_file(zip_path)
    _safe_remove_dir(tmp_dir)


def _safe_remove_file(path: Path) -> None:
    try:
        path.unlink(missing_ok=True)
    except Exception:  # noqa: BLE001
        pass


def _safe_remove_dir(path: Path) -> None:
    shutil.rmtree(path, ignore_errors=True)
