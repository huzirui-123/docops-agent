"""Subprocess runner for API execution isolation."""

from __future__ import annotations

import json
import os
import time
from dataclasses import dataclass
from multiprocessing.connection import Connection
from pathlib import Path
from typing import Literal

from docx import Document

from apps.cli.io import (
    build_output_paths,
    write_render_output_atomic,
    write_suggested_policy_atomic,
)
from core.format.policy_loader import load_policy
from core.format.suggested_policy import build_suggested_policy
from core.orchestrator.pipeline import run_task
from core.skills.base import Skill
from core.skills.models import TaskSpec
from core.skills.registry import create_skill
from core.utils.errors import MissingRequiredFieldsError, TemplateError

FormatMode = Literal["report", "strict", "off"]
FormatBaseline = Literal["template", "policy"]
FormatFixMode = Literal["none", "safe"]


@dataclass(frozen=True)
class RunnerRequest:
    """Serializable runner request payload."""

    tmp_dir: str
    template_path: str
    task_path: str
    skill_name: str
    policy_path: str | None
    unsupported_mode: Literal["error", "warn"]
    format_mode: FormatMode
    format_baseline: FormatBaseline
    format_fix_mode: FormatFixMode
    export_suggested_policy: bool


@dataclass(frozen=True)
class RunnerResponse:
    """Serializable runner response payload."""

    ok: bool
    exit_code: int | None
    message: str
    error_type: str | None = None
    error_message: str | None = None


def run_pipeline_worker(request: RunnerRequest, send_conn: Connection) -> None:
    """Run one docops job inside a subprocess and return a structured response."""

    try:
        _apply_test_hooks(Path(request.tmp_dir))
        response = _run_worker_request(request)
    except Exception as exc:  # noqa: BLE001
        response = RunnerResponse(
            ok=False,
            exit_code=None,
            message="runner failed",
            error_type=exc.__class__.__name__,
            error_message=str(exc),
        )

    try:
        send_conn.send(response)
    finally:
        send_conn.close()


def _run_worker_request(request: RunnerRequest) -> RunnerResponse:
    tmp_dir = Path(request.tmp_dir)
    template_path = Path(request.template_path)
    task_path = Path(request.task_path)
    policy_path = Path(request.policy_path) if request.policy_path is not None else None

    task_spec = _load_task_spec(task_path)
    skill = _resolve_skill(request.skill_name)
    policy = load_policy(policy_path)

    # Keep a pristine template copy for suggested policy generation.
    template_for_suggested = (
        Document(str(template_path)) if request.export_suggested_policy else None
    )
    template_document = Document(str(template_path))

    try:
        output = run_task(
            task_spec=task_spec,
            template_document=template_document,
            skill=skill,
            policy=policy,
            unsupported_mode=request.unsupported_mode,
            format_mode=request.format_mode,
            format_baseline=request.format_baseline,
            format_fix_mode=request.format_fix_mode,
        )
        exit_code = 0
        message = "success"
    except TemplateError as exc:
        if exc.render_output is None:
            raise RuntimeError("TemplateError missing render_output") from exc
        output = exc.render_output
        exit_code = 3
        message = "template unsupported placeholders"
    except MissingRequiredFieldsError as exc:
        if exc.render_output is None:
            raise RuntimeError("MissingRequiredFieldsError missing render_output") from exc
        output = exc.render_output
        exit_code = 2
        message = "missing required fields"

    if request.format_mode == "strict" and not output.format_report.passed:
        exit_code = 4
        message = "format validation failed"

    output_paths = build_output_paths(tmp_dir)
    write_render_output_atomic(output_paths, output)

    if request.export_suggested_policy:
        if template_for_suggested is None:
            raise RuntimeError("Template document is required for suggested policy export")
        suggested_payload = build_suggested_policy(template_for_suggested, policy)
        write_suggested_policy_atomic(tmp_dir / "out.suggested_policy.yaml", suggested_payload)

    return RunnerResponse(ok=True, exit_code=exit_code, message=message)


def _resolve_skill(skill_name: str) -> Skill:
    return create_skill(skill_name)


def _load_task_spec(task_path: Path) -> TaskSpec:
    raw = json.loads(task_path.read_text(encoding="utf-8"))
    if not isinstance(raw, dict):
        raise ValueError("task JSON must be an object")
    return TaskSpec.model_validate(raw)


def _apply_test_hooks(tmp_dir: Path) -> None:
    if os.getenv("DOCOPS_TEST_MODE") == "1":
        (tmp_dir / "runner.pid").write_text(str(os.getpid()), encoding="utf-8")

    sleep_seconds = os.getenv("DOCOPS_TEST_SLEEP_SECONDS")
    if sleep_seconds:
        time.sleep(float(sleep_seconds))
