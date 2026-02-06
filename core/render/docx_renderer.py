"""Docx renderer for run-level placeholder replacement."""

from __future__ import annotations

from collections import defaultdict
from collections.abc import Mapping
from typing import Literal

from docx.document import Document as DocxDocument
from docx.text.run import Run

from core.format.run_ids import iter_paragraph_run_contexts
from core.render.models import (
    MissingFieldsReport,
    RenderOutput,
    ReplaceLogEntry,
    ReplaceReport,
    ReplaceSummary,
    empty_format_report,
)
from core.skills.models import SkillResult
from core.templates.models import Occurrence, ParseResult
from core.templates.placeholder_parser import parse_placeholders
from core.utils.errors import TemplateError


def render_docx(
    document: DocxDocument,
    skill_result: SkillResult,
    unsupported_mode: Literal["error", "warn"] = "error",
) -> RenderOutput:
    """Render placeholders in a docx document with run-level safety constraints."""

    if unsupported_mode not in {"error", "warn"}:
        raise ValueError(f"Unsupported mode: {unsupported_mode}")

    parse_result = parse_placeholders(document, strict=False)
    template_fields = list(parse_result.fields)

    entries: list[ReplaceLogEntry] = []
    touched_runs: set[str] = set()

    for item in parse_result.unsupported:
        entries.append(
            ReplaceLogEntry(
                status="unsupported",
                run_id=item.run_id,
                paragraph_path=_paragraph_path_from_run_id(item.run_id),
                start=item.start,
                end=item.end,
                original_text=item.text,
                reason=item.kind,
            )
        )

    unsupported_count = len(parse_result.unsupported)
    missing_fields_report = _compute_missing_fields(template_fields, skill_result)

    if unsupported_count > 0 and unsupported_mode == "error":
        replace_report = _build_replace_report(
            parse_result=parse_result,
            entries=entries,
            touched_runs=touched_runs,
            unsupported_mode=unsupported_mode,
        )
        partial_output = RenderOutput(
            document=document,
            parse_result=parse_result,
            template_fields=template_fields,
            replace_report=replace_report,
            missing_fields=missing_fields_report,
            format_report=empty_format_report(),
        )
        raise TemplateError(
            "Unsupported placeholders found in template",
            result=parse_result,
            replace_report=replace_report,
            render_output=partial_output,
        )

    run_lookup = _build_run_lookup(document)
    replaced_entries, touched_runs = _replace_supported_occurrences(
        occurrences=parse_result.occurrences,
        run_lookup=run_lookup,
        field_values=skill_result.field_values,
        initial_touched=touched_runs,
    )
    entries.extend(replaced_entries)

    replace_report = _build_replace_report(
        parse_result=parse_result,
        entries=entries,
        touched_runs=touched_runs,
        unsupported_mode=unsupported_mode,
    )

    return RenderOutput(
        document=document,
        parse_result=parse_result,
        template_fields=template_fields,
        replace_report=replace_report,
        missing_fields=missing_fields_report,
        format_report=empty_format_report(),
    )


def _replace_supported_occurrences(
    occurrences: list[Occurrence],
    run_lookup: Mapping[str, Run],
    field_values: dict[str, str],
    initial_touched: set[str],
) -> tuple[list[ReplaceLogEntry], set[str]]:
    grouped: dict[str, list[Occurrence]] = defaultdict(list)
    for occurrence in occurrences:
        grouped[occurrence.run_id].append(occurrence)

    entries: list[ReplaceLogEntry] = []
    touched_runs = set(initial_touched)

    for run_id in sorted(grouped.keys()):
        run = run_lookup.get(run_id)
        if run is None:
            continue

        run_text = run.text or ""
        for occurrence in sorted(grouped[run_id], key=lambda item: item.start, reverse=True):
            token_text = run_text[occurrence.start : occurrence.end]
            replacement = field_values.get(occurrence.field_name)
            if replacement is None:
                entries.append(
                    ReplaceLogEntry(
                        status="missing",
                        field_name=occurrence.field_name,
                        run_id=run_id,
                        paragraph_path=_paragraph_path_from_run_id(run_id),
                        start=occurrence.start,
                        end=occurrence.end,
                        original_text=token_text,
                        reason="missing_field",
                    )
                )
                continue

            run_text = (
                run_text[: occurrence.start] + replacement + run_text[occurrence.end :]
            )
            touched_runs.add(run_id)
            entries.append(
                ReplaceLogEntry(
                    status="replaced",
                    field_name=occurrence.field_name,
                    run_id=run_id,
                    paragraph_path=_paragraph_path_from_run_id(run_id),
                    start=occurrence.start,
                    end=occurrence.end,
                    original_text=token_text,
                    new_text=replacement,
                )
            )

        run.text = run_text

    return entries, touched_runs


def _build_run_lookup(document: DocxDocument) -> dict[str, Run]:
    run_lookup: dict[str, Run] = {}
    for context in iter_paragraph_run_contexts(document, include_tables=True):
        for run_id, run in zip(context.run_ids, context.paragraph.runs, strict=False):
            run_lookup[run_id] = run
    return run_lookup


def _compute_missing_fields(
    template_fields: list[str], skill_result: SkillResult
) -> MissingFieldsReport:
    template_field_set = set(template_fields)
    provided = set(skill_result.field_values.keys())
    missing_in_template = template_field_set - provided

    required = set(skill_result.required_fields)
    optional = set(skill_result.optional_fields)

    missing_required = sorted(missing_in_template & required)
    missing_optional = (missing_in_template & optional) | (
        missing_in_template - (required | optional)
    )

    return MissingFieldsReport(
        missing_required=missing_required,
        missing_optional=sorted(missing_optional),
    )


def _build_replace_report(
    parse_result: ParseResult,
    entries: list[ReplaceLogEntry],
    touched_runs: set[str],
    unsupported_mode: Literal["error", "warn"],
) -> ReplaceReport:
    replaced_count = sum(1 for item in entries if item.status == "replaced")
    missing_count = sum(1 for item in entries if item.status == "missing")
    unsupported_count = sum(1 for item in entries if item.status == "unsupported")

    summary = ReplaceSummary(
        total_placeholders=len(parse_result.occurrences) + len(parse_result.unsupported),
        replaced_count=replaced_count,
        missing_count=missing_count,
        had_unsupported=unsupported_count > 0,
        unsupported_count=unsupported_count,
        unsupported_mode=unsupported_mode,
    )
    return ReplaceReport(entries=entries, summary=summary, touched_runs=sorted(touched_runs))


def _paragraph_path_from_run_id(run_id: str | None) -> str | None:
    if run_id is None:
        return None
    return run_id.split(":r", maxsplit=1)[0]
